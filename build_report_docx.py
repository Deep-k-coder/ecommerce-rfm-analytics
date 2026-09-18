"""
build_report_docx.py
Generates the publication-grade Microsoft Word (.docx) project report:
Deep_Koshiya_ProjectReport.docx
"""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets the background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def create_document():
    base_dir = Path(__file__).resolve().parent
    img_dir = base_dir / "report_images"
    doc = Document()

    # Configure Standard Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Colors
    NAVY = RGBColor(15, 23, 42)      # Primary text #0f172a
    INDIGO = RGBColor(79, 70, 229)   # Accent #4f46e5
    SLATE = RGBColor(100, 116, 139)  # Secondary #64748b

    # -------------------------------------------------------------
    # COVER / HEADER TITLE BLOCK
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("E-COMMERCE CUSTOMER LIFECYCLE &\nRFM SEGMENTATION DASHBOARD")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("End-to-End Analytics Engineering, Algorithmic Clustering & Strategic CRM Activation")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = INDIGO

    # Metadata Summary Box
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Candidate / Author", "Deep Koshiya"),
        ("Project Scope", "Customer Lifecycle Analytics, RFM Scoring, K-Means Clustering & BI Engineering"),
        ("Technology Stack", "Python 3.9+, Pandas, Scikit-Learn, SciPy, Matplotlib, HTML5/Tailwind/Chart.js"),
        ("Submission Date", "September 2026")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.text = k
        c1.text = v
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(9.5)
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Helper function for styled section headings
    def add_section_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(15)
            run.font.color.rgb = INDIGO
        elif level == 2:
            run.font.size = Pt(12.5)
            run.font.color.rgb = NAVY
        return p

    def add_body_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.bold = True
            r_pre.font.size = Pt(10)
        r_main = p.add_run(text)
        r_main.font.size = Pt(10)
        return p

    # -------------------------------------------------------------
    # 1. EXECUTIVE SUMMARY
    # -------------------------------------------------------------
    add_section_heading("1. Executive Summary")
    add_body_p(
        "In modern multi-category digital commerce, treating an entire customer base as a homogenous entity leads to suboptimal marketing spend, inflated Customer Acquisition Costs (CAC), and customer churn blindspots. This comprehensive project establishes an end-to-end data analytics engineering pipeline that takes raw transactional event logs and refines them into actionable behavioral cohorts."
    )
    add_body_p(
        "By synthesizing data cleaning best practices (including Tukey's IQR outlier isolation and credit note segregation), customer-level RFM feature engineering, unsupervised machine learning (K-Means), and heuristic quantile scoring (pd.qcut), the system generates an enterprise star-schema data mart and an interactive executive BI dashboard. Key commercial findings reveal a strong Pareto concentration: 27.4% of customers ('Champions') generate 72.5% of gross revenues, while $2,384+ in high-margin spend is vulnerable within at-risk, dormant cohorts."
    )

    # -------------------------------------------------------------
    # 2. BUSINESS PROBLEM STATEMENT & DOMAIN CONTEXT
    # -------------------------------------------------------------
    add_section_heading("2. Business Problem & Commercial Context")
    add_body_p(
        "Online retailers face distinct commercial challenges across customer acquisition, retention, and loyalty development:"
    )
    add_body_p("Marketing campaigns delivered uniformly dilute returns on ad spend (ROAS). High-value accounts receive standard discount coupons they do not need, while churn-risk customers are neglected until they lapse permanently.", "• Budget Misallocation: ")
    add_body_p("Without longitudinal entity tracking, management cannot distinguish between a dip in seasonal demand versus a systemic churn trend among top-tier accounts.", "• Churn Vulnerability: ")
    add_body_p("Unfiltered transactional datasets often contain B2B distributor bulk orders, cancelled credit invoices, and unauthenticated guest sessions, creating artificial distortions in lifetime value models.", "• Data Integrity Impediments: ")

    # -------------------------------------------------------------
    # 3. DATA REFINERY & QUALITY PIPELINE
    # -------------------------------------------------------------
    add_section_heading("3. Data Cleaning & Refinery Pipeline")
    add_body_p(
        "Raw transactional logs conform to the standard UCI Online Retail format with 7 initial attributes: InvoiceNo, StockCode, Description, Quantity, InvoiceDate, UnitPrice, and CustomerID. The cleaning architecture implements four defensive stages:"
    )
    add_body_p("Casting string identifiers, enforcing ISO-8601 timestamps, and trimming extraneous whitespace across SKU descriptions.", "1. Schema Enforcement & Typing: ")
    add_body_p("Transactions lacking CustomerID represent anonymous guest checkouts. Because they cannot be linked across multiple sessions, these 1,888 records are dropped to maintain longitudinal customer profiles.", "2. Entity Identification: ")
    add_body_p("Invoices starting with 'C' or containing negative quantities/zero prices represent product returns, damaged items, or system adjustments. Rather than discarding them entirely, 713 records were segregated into fact_reversals_audit.csv for accounting reconciliation.", "3. Reverse Logistics Isolation: ")
    add_body_p("Bulk wholesale orders distort consumer clustering. Using Tukey's Interquartile Range method (IQR = Q3 - Q1), fences were constructed to bound quantities within [1, 24] and prices within (0, 7.50].", "4. Statistical Outlier Removal: ")

    # -------------------------------------------------------------
    # 4. FEATURE ENGINEERING (RFM METRIC DERIVATION)
    # -------------------------------------------------------------
    add_section_heading("4. Feature Engineering: RFM Derivation")
    add_body_p(
        "The cleaned transactional stream was aggregated into a customer-level analytical base table (ABT) using a fixed operational snapshot date: T_obs = max(InvoiceDate) + 1 day (December 11, 2023)."
    )
    add_body_p("Recency (R) = Time elapsed in integer days between the customer's most recent transaction and T_obs. Lower values indicate higher current engagement.", "• Recency (R): ")
    add_body_p("Frequency (F) = Distinct count of unique invoices (nunique) placed by the customer. Evaluating distinct checkouts prevents multi-item single carts from falsely appearing as repeat visits.", "• Frequency (F): ")
    add_body_p("Monetary (M) = Cumulative net spend across all valid items (Quantity * UnitPrice).", "• Monetary (M): ")
    add_body_p("Average Order Value (AOV) = Monetary / Frequency. A secondary metric indicating basket efficiency.", "• Average Order Value (AOV): ")

    # -------------------------------------------------------------
    # 5. MACHINE LEARNING SEGMENTATION & CLUSTERING
    # -------------------------------------------------------------
    add_section_heading("5. Unsupervised Machine Learning (K-Means)")
    add_body_p(
        "Due to the heavy positive skew inherent to consumer spend, feeding raw RFM values directly into K-Means causes cluster centroids to be dominated by extreme spenders. The pipeline executes:"
    )
    add_body_p("Compresses long right tails into normal approximations, handling low values smoothly.", "1. Log1p Transformation: ")
    add_body_p("Standardizes features to zero mean and unit variance (z = (x - mu) / sigma), ensuring equal geometric weighting in Euclidean distance calculations.", "2. Z-Score Standardization: ")
    add_body_p("Evaluated cluster numbers k in [2..6] using Within-Cluster Sum of Squares (Inertia) and Silhouette Scores. Optimal separation occurred at k=4 (Silhouette Score: 0.3817, Inertia: 430.5).", "3. Diagnostics & Cluster Tuning: ")

    # Embed K-Means Diagnostics Image
    diag_img = img_dir / "fig_kmeans_diagnostics.png"
    if diag_img.exists():
        doc.add_paragraph().paragraph_format.space_before = Pt(6)
        doc.add_picture(str(diag_img), width=Inches(6.0))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 1: K-Means Elbow Curve (Inertia) and Silhouette Score Diagnostics across Cluster Counts (k=2 to 6).")
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = SLATE

    # -------------------------------------------------------------
    # 6. QUANTILE RFM SCORING & OPERATIONAL SEGMENTS
    # -------------------------------------------------------------
    add_section_heading("6. Quantile RFM Scoring & Taxonomy")
    add_body_p(
        "To provide operational consistency for marketing automation, each customer was also assigned 1–5 relative scores using quintile binning (pd.qcut):"
    )
    add_body_p("Score 5 represents the most recent buyers (lowest days); Score 1 represents longest inactivity.", "• Recency Score (R_Score): ")
    add_body_p("Score 5 represents top repeat buyers; Score 1 represents single purchase shoppers.", "• Frequency Score (F_Score): ")
    add_body_p("Score 5 represents top gross spenders; Score 1 represents lowest monetary contribution.", "• Monetary Score (M_Score): ")
    add_body_p("A 3-digit string (e.g., '555' for top champions, '111' for dormant single purchasers).", "• RFM Composite: ")

    # Embed Spatial Scatter & Revenue Share
    scatter_img = img_dir / "fig_spatial_scatter.png"
    if scatter_img.exists():
        doc.add_paragraph().paragraph_format.space_before = Pt(6)
        doc.add_picture(str(scatter_img), width=Inches(6.0))
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap2 = p_cap2.add_run("Figure 2: Spatial Customer Migration Map (Recency vs. Spend, Bubble Size = Frequency).")
        r_cap2.font.size = Pt(8.5)
        r_cap2.font.italic = True
        r_cap2.font.color.rgb = SLATE

    # -------------------------------------------------------------
    # 7. EMPIRICAL FINDINGS & PORTFOLIO BREAKDOWN
    # -------------------------------------------------------------
    add_section_heading("7. Portfolio Health & Segment Analytics")
    add_body_p(
        "Across the portfolio of 540 active customers generating $65,132.81 in total spend, behavioral segmentation yielded 8 operational cohorts:"
    )

    # Segment Performance Table
    table = doc.add_table(rows=9, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Business Segment", "Users", "User %", "Total Spend", "Spend %", "Avg R (d)", "Avg F"]
    hdr_row = table.rows[0]
    for j, h in enumerate(headers):
        cell = hdr_row.cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)
        set_cell_background(cell, "0F172A")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    table_data = [
        ("Champions", "148", "27.4%", "$47,250.31", "72.5%", "13.6d", "18.4"),
        ("Loyal Customers", "95", "17.6%", "$6,346.33", "9.7%", "37.7d", "3.8"),
        ("Hibernating", "116", "21.5%", "$3,745.34", "5.8%", "139.7d", "1.9"),
        ("Mid-Market Regulars", "13", "2.4%", "$2,508.67", "3.9%", "44.5d", "11.3"),
        ("At Risk (High Value)", "23", "4.3%", "$2,384.72", "3.7%", "92.8d", "5.9"),
        ("Lost / Inactive", "77", "14.3%", "$1,255.78", "1.9%", "163.0d", "1.0"),
        ("Needs Attention", "33", "6.1%", "$846.77", "1.3%", "47.2d", "1.5"),
        ("Recent Converts", "35", "6.5%", "$794.89", "1.2%", "15.9d", "1.3")
    ]
    for i, row_vals in enumerate(table_data):
        row = table.rows[i + 1]
        for j, val in enumerate(row_vals):
            cell = row.cells[j]
            cell.text = val
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            if i % 2 == 0:
                set_cell_background(cell, "F8FAFC")
            else:
                set_cell_background(cell, "FFFFFF")

    doc.add_paragraph().paragraph_format.space_before = Pt(8)

    # -------------------------------------------------------------
    # 8. STRATEGIC CRM ACTION PLAYBOOK
    # -------------------------------------------------------------
    add_section_heading("8. Strategic CRM Action Playbooks")
    add_body_p("Deploy VIP concierge treatment, early product releases, and non-discount rewards (loyalty gifts). Do not over-discount, as this segment exhibits low price sensitivity.", "• Champions (72.5% Rev): ")
    add_body_p("Introduce category expansion programs, bundle recommendations, and threshold loyalty point boosters (e.g., 'Spend $50 to earn 2x points').", "• Loyal Customers (9.7% Rev): ")
    add_body_p("Urgent win-back required. Deploy personalized multi-channel outreach (SMS/email) with aggressive time-limited vouchers ($15 off $50) before permanent churn.", "• At Risk - High Value ($2,384 exposed): ")
    add_body_p("Deploy automated onboarding drip campaigns. Provide an immediate 10% coupon valid for 7 days to incentivize second order conversion.", "• Recent Converts / Promising: ")
    add_body_p("Suppress completely from paid acquisition retargeting lists (Google Ads, Meta Ads) to eliminate budget waste. Retain only low-cost automated email cycles.", "• Lost / Inactive: ")

    # -------------------------------------------------------------
    # 9. EXECUTIVE BI DASHBOARD ARCHITECTURE
    # -------------------------------------------------------------
    add_section_heading("9. Executive BI Dashboard Architecture")
    add_body_p(
        "To enable self-service exploration by marketing executives and commercial directors, the pipeline compiles a responsive executive web dashboard and exports clean star-schema datasets:"
    )
    add_body_p("DIM_CUSTOMER_RFM (CustomerID PK, Recency, Frequency, Monetary, AOV, R_Score, F_Score, M_Score, Business_Segment, Recommended_Action) joined to FACT_TRANSACTIONS (InvoiceNo, StockCode, Quantity, UnitPrice, LineItemTotal).", "• Star-Schema Model: ")
    add_body_p("1920x1080 FHD executive layout featuring 6 KPI scorecards, segment donut composition, spatial migration scatterplot, K-Means diagnostics, and a searchable activation table.", "• Interactive Dashboard: ")

    # -------------------------------------------------------------
    # 10. CONCLUSION & FUTURE ENHANCEMENTS
    # -------------------------------------------------------------
    add_section_heading("10. Conclusion & Future Roadmap")
    add_body_p(
        "This project establishes a verified, production-grade analytics framework that translates raw e-commerce logs into distinct, actionable customer segments. By pairing statistical outlier cleaning with both machine learning (K-Means) and rule-based quantile scoring, the organization gains clear operational visibility into revenue concentration and churn vulnerabilities."
    )
    add_body_p("Integrating Buy 'Til You Die (BTYD) probabilistic models (BG/NBD and Gamma-Gamma) to predict individual customer transaction probabilities and forward-looking 12-month expected monetary value.", "• Probabilistic CLV Modeling: ")
    add_body_p("Triggering automated webhook payloads into CRM platforms (Klaviyo, Braze, Salesforce) based on daily cohort transitions.", "• Automated Event-Driven Webhooks: ")

    out_docx = base_dir / "Deep_Koshiya_ProjectReport.docx"
    doc.save(str(out_docx))
    print(f"[REPORT] Successfully generated comprehensive Word project report at {out_docx}")

if __name__ == "__main__":
    create_document()
