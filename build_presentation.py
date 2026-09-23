"""
build_presentation.py
Generates the definitive 10-slide PowerPoint presentation (.pptx) for the
First Progress Work Project Report (24/09/2026):
- Degree: B.Sc. in Artificial Intelligence & Data Science, Semester V (5th Semester)
- Institution: Institute of Advanced Research (IAR), Gandhinagar
- NO DEPARTMENT NAME mentioned anywhere in the presentation.
- Official IAR banner at top of Title & Thank You slides.
- 16:9 Widescreen format (13.333" x 7.5").
- Exactly 10 slides, culminating in an elegant "Thank You & Q&A" slide.
- High-res UML & DFD diagrams (Level 0, 1, 2 DFDs and Use Case diagram).
- Color Palette: IAR Maroon (#7A003C), Deep Navy (#1E1B4B), Slate (#475569), Clean Whites & Light Grays.
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Brand Colors
COLOR_MAROON = RGBColor(122, 0, 60)      # IAR Brand Maroon #7A003C
COLOR_NAVY = RGBColor(30, 27, 75)        # Deep Navy #1E1B4B
COLOR_DARK = RGBColor(15, 23, 42)        # Body Dark #0F172A
COLOR_SLATE = RGBColor(71, 85, 105)      # Slate Gray #475569
COLOR_MUTED = RGBColor(148, 163, 184)    # Light Slate #94A3B8
COLOR_CARD_BG = RGBColor(248, 250, 252)  # Card Background #F8FAFC
COLOR_BORDER = RGBColor(226, 232, 240)   # Border #E2E8F0
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_EMERALD = RGBColor(16, 185, 129)   # Accent Green #10B981
COLOR_BLUE = RGBColor(37, 99, 235)       # Accent Blue #2563EB
COLOR_AMBER = RGBColor(217, 119, 6)      # Accent Amber #D97706

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    base_dir = Path("/Users/deep/.gemini/antigravity/scratch/ecommerce_rfm_analytics")
    img_dir = base_dir / "report_images"

    def add_header(slide, category_text, title_text, slide_num):
        # Top Maroon accent line
        accent_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08)
        )
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = COLOR_MAROON
        accent_line.line.color.rgb = COLOR_MAROON

        # Category Chip
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.25), Inches(8.5), Inches(0.35))
        p_cat = tb_cat.text_frame.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = 'Arial'
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_MAROON

        # Slide Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(9.5), Inches(0.65))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = 'Arial'
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_NAVY

        # Slide Number Badge
        tb_num = slide.shapes.add_textbox(Inches(11.2), Inches(0.35), Inches(1.3), Inches(0.4))
        p_num = tb_num.text_frame.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = f"{slide_num:02d} / 10"
        p_num.font.name = 'Arial'
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_SLATE

        # Subtle divider
        div = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.22), Inches(11.733), Inches(0.015)
        )
        div.fill.solid()
        div.fill.fore_color.rgb = COLOR_BORDER
        div.line.color.rgb = COLOR_BORDER

    def add_footer(slide):
        tb_foot = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
        p_foot = tb_foot.text_frame.paragraphs[0]
        p_foot.text = "B.Sc. in Artificial Intelligence & Data Science (Sem V)  •  Institute of Advanced Research  •  First Progress Work (24/09/2026)"
        p_foot.font.name = 'Arial'
        p_foot.font.size = Pt(9)
        p_foot.font.color.rgb = COLOR_MUTED

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # =========================================================================
    # SLIDE 1: COVER SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)

    # Top & Bottom accent bands
    top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_MAROON
    top_bar.line.color.rgb = COLOR_MAROON

    bot_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.38), Inches(13.333), Inches(0.12))
    bot_bar.fill.solid()
    bot_bar.fill.fore_color.rgb = COLOR_NAVY
    bot_bar.line.color.rgb = COLOR_NAVY

    # IAR Banner
    iar_banner = img_dir / "iar_banner_hd.png"
    if iar_banner.exists():
        slide1.shapes.add_picture(str(iar_banner), Inches(3.9), Inches(0.45), width=Inches(5.5))

    # University & Program Heading (NO DEPARTMENT NAME)
    tb_univ = slide1.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(11.333), Inches(0.6))
    p_u = tb_univ.text_frame.paragraphs[0]
    p_u.alignment = PP_ALIGN.CENTER
    p_u.text = "INSTITUTE OF ADVANCED RESEARCH (IAR)"
    p_u.font.name = 'Arial'
    p_u.font.size = Pt(17)
    p_u.font.bold = True
    p_u.font.color.rgb = COLOR_MAROON

    p_u2 = tb_univ.text_frame.add_paragraph()
    p_u2.alignment = PP_ALIGN.CENTER
    p_u2.text = "The University for Innovation  |  Gandhinagar, Gujarat"
    p_u2.font.name = 'Arial'
    p_u2.font.size = Pt(11)
    p_u2.font.italic = True
    p_u2.font.color.rgb = COLOR_SLATE

    # Milestone Badge Card
    badge = add_card(slide1, Inches(3.2), Inches(2.35), Inches(6.933), Inches(0.42), bg_color=RGBColor(239, 246, 255), border_color=COLOR_BLUE)
    tb_badge = slide1.shapes.add_textbox(Inches(3.2), Inches(2.35), Inches(6.933), Inches(0.42))
    p_b = tb_badge.text_frame.paragraphs[0]
    p_b.alignment = PP_ALIGN.CENTER
    p_b.text = "FIRST PROGRESS WORK PRESENTATION  •  24TH SEPTEMBER 2026  •  10 MARKS"
    p_b.font.name = 'Arial'
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_BLUE

    # Main Project Title
    tb_t = slide1.shapes.add_textbox(Inches(1.0), Inches(2.9), Inches(11.333), Inches(1.3))
    p_t = tb_t.text_frame.paragraphs[0]
    p_t.alignment = PP_ALIGN.CENTER
    p_t.text = "E-COMMERCE CUSTOMER LIFECYCLE &\nRFM SEGMENTATION DASHBOARD"
    p_t.font.name = 'Arial'
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_NAVY

    p_sub = tb_t.text_frame.add_paragraph()
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.text = "An Applied Machine Learning & Analytics Engineering System for Behavioral Cohort Optimization"
    p_sub.font.name = 'Arial'
    p_sub.font.size = Pt(13)
    p_sub.font.italic = True
    p_sub.font.color.rgb = COLOR_MAROON

    # Metadata Cards (Scholar & Academic Context)
    c1 = add_card(slide1, Inches(1.5), Inches(4.5), Inches(5.0), Inches(2.3))
    tb_c1 = slide1.shapes.add_textbox(Inches(1.7), Inches(4.65), Inches(4.6), Inches(2.0))
    tf1 = tb_c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "PROJECT SCHOLAR"
    p.font.name = 'Arial'
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    lines_c1 = [
        ("Candidate Name: ", "Mr. Deep Koshiya"),
        ("Program: ", "B.Sc. in AI & Data Science"),
        ("Semester: ", "Semester V (5th Semester)"),
        ("Academic Year: ", "2026 – 2027"),
        ("Course: ", "Project - 1 (CE-502 / PRJ-701)")
    ]
    for lbl, val in lines_c1:
        p = tf1.add_paragraph()
        r1 = p.add_run()
        r1.text = lbl
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_DARK
        r2 = p.add_run()
        r2.text = val
        r2.font.size = Pt(10)
        r2.font.color.rgb = COLOR_DARK

    c2 = add_card(slide1, Inches(6.8), Inches(4.5), Inches(5.0), Inches(2.3))
    tb_c2 = slide1.shapes.add_textbox(Inches(7.0), Inches(4.65), Inches(4.6), Inches(2.0))
    tf2 = tb_c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "EVALUATION & MILESTONE SCOPE"
    p.font.name = 'Arial'
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    lines_c2 = [
        ("Milestone: ", "First Progress Work Evaluation"),
        ("Evaluation Date: ", "24th September 2026 (10 AM - 5 PM)"),
        ("Marks Allotted: ", "10 Marks (Curriculum Deliverable)"),
        ("Syllabus Scope: ", "Sections 1.1, 1.2, 1.3, 2.1, 2.2"),
        ("Supervision: ", "Assigned Project Supervisor (IAR)")
    ]
    for lbl, val in lines_c2:
        p = tf2.add_paragraph()
        r1 = p.add_run()
        r1.text = lbl
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_DARK
        r2 = p.add_run()
        r2.text = val
        r2.font.size = Pt(10)
        r2.font.color.rgb = COLOR_DARK

    # =========================================================================
    # SLIDE 2: 1.1 PROBLEM DEFINITION & COMMERCIAL CONTEXT
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Section 1.1  •  About the System", "Problem Definition & Commercial Context", 2)
    add_footer(slide2)

    col_w = Inches(3.64)
    gap = Inches(0.4)
    y_top = Inches(1.5)
    card_h = Inches(5.2)

    # Card 1: The Commercial Dilemma
    add_card(slide2, Inches(0.8), y_top, col_w, card_h)
    tb = slide2.shapes.add_textbox(Inches(0.95), y_top + Inches(0.15), col_w - Inches(0.3), card_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. Commercial Retention Dilemma"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    bullets1 = [
        ("Surging Acquisition Costs: ", "Digital advertising CAC has increased over 60%, making organic customer retention the primary driver of profitability."),
        ("Homogenous Marketing: ", "Treating all shoppers equally leads to margin-eroding discounts for customers who would have purchased anyway."),
        ("Undetected Churn Risk: ", "Historically high-value accounts slip into dormancy silently without early churn detection triggers."),
        ("Missed Lifetime Value: ", "E-commerce retailers fail to identify and nurture accounts transitioning into Champions.")
    ]
    for b_title, b_desc in bullets1:
        p = tf.add_paragraph()
        p.space_before = Pt(8)
        r0 = p.add_run()
        r0.text = "• " + b_title
        r0.font.bold = True
        r0.font.size = Pt(10.5)
        r0.font.color.rgb = COLOR_NAVY
        r1 = p.add_run()
        r1.text = b_desc
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_DARK

    # Card 2: OLTP System Limitations
    add_card(slide2, Inches(0.8) + col_w + gap, y_top, col_w, card_h)
    tb = slide2.shapes.add_textbox(Inches(0.95) + col_w + gap, y_top + Inches(0.15), col_w - Inches(0.3), card_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. Limitations of OLTP Systems"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    bullets2 = [
        ("Order-Centric Storage: ", "Traditional databases record individual shopping transactions, not longitudinal customer behavioral lifecycles."),
        ("Data Quality Deficits: ", "Event logs contain guest checkouts lacking CustomerID, cancellations, and administrative line-items."),
        ("Distribution Anomalies: ", "Extreme wholesale B2B distributor purchases contaminate standard consumer metrics."),
        ("No Analytical Layer: ", "Decision-makers lack an automated transformation pipeline bridging raw order rows with strategic marketing actions.")
    ]
    for b_title, b_desc in bullets2:
        p = tf.add_paragraph()
        p.space_before = Pt(8)
        r0 = p.add_run()
        r0.text = "• " + b_title
        r0.font.bold = True
        r0.font.size = Pt(10.5)
        r0.font.color.rgb = COLOR_NAVY
        r1 = p.add_run()
        r1.text = b_desc
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_DARK

    # Card 3: Targeted System Solution
    add_card(slide2, Inches(0.8) + (col_w + gap) * 2, y_top, col_w, card_h)
    tb = slide2.shapes.add_textbox(Inches(0.95) + (col_w + gap) * 2, y_top + Inches(0.15), col_w - Inches(0.3), card_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "3. Targeted Project Solution"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    bullets3 = [
        ("Automated Refinery: ", "Cleanses raw transactions, isolates reversals into an audit ledger, and caps outliers via Tukey IQR fences."),
        ("Vectorized RFM Feature Store: ", "Aggregates transactions into Recency, Frequency, Monetary spend, and Average Order Value (AOV)."),
        ("Skew-Resistant ML Clustering: ", "Applies log1p and StandardScaler before K-Means modeling, validated via Silhouette diagnostics."),
        ("Executive BI Dashboard: ", "Interactive web UI with portfolio KPI scorecards, spatial scatterplots, and exportable CRM cohorts.")
    ]
    for b_title, b_desc in bullets3:
        p = tf.add_paragraph()
        p.space_before = Pt(8)
        r0 = p.add_run()
        r0.text = "• " + b_title
        r0.font.bold = True
        r0.font.size = Pt(10.5)
        r0.font.color.rgb = COLOR_NAVY
        r1 = p.add_run()
        r1.text = b_desc
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_DARK

    # =========================================================================
    # SLIDE 3: 1.2 REQUIREMENT SPECIFICATIONS
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Section 1.2  •  Product Specifications", "Requirement Specifications (Functional & Non-Functional)", 3)
    add_footer(slide3)

    # Left Column: Functional Requirements (Width 6.8")
    add_card(slide3, Inches(0.8), Inches(1.5), Inches(6.8), Inches(5.2))
    tb_fr = slide3.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(6.4), Inches(4.9))
    tf_fr = tb_fr.text_frame
    tf_fr.word_wrap = True
    p = tf_fr.paragraphs[0]
    p.text = "Functional Requirements Specifications (FR-01 to FR-06)"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    fr_items = [
        ("FR-01: Ingestion & Schema Typing: ", "Ingests raw CSV streams, validates typing (InvoiceDate to datetime64[ns], Quantity/UnitPrice to floats), and sanitizes SKU descriptions."),
        ("FR-02: Missing Entity Resolution: ", "Detects and filters guest checkouts missing CustomerID to protect longitudinal cohort integrity."),
        ("FR-03: Reverse Logistics Isolation: ", "Routes cancellations ('^C') and credit adjustments to an isolated audit store (fact_reversals_audit.csv) without distorting purchase counts."),
        ("FR-04: Tukey IQR Outlier Capping: ", "Applies fences [Q1 - 1.5*IQR, Q3 + 1.5*IQR] on Quantity and Price, neutralizing bulk distributor volume spikes."),
        ("FR-05: Vectorized RFM Derivation: ", "Computes Recency days relative to T_obs = max(Date) + 1, Frequency (distinct visits), and cumulative Monetary spend."),
        ("FR-06: ML Clustering & Scoring: ", "Standardizes features via log1p and StandardScaler, fits K-Means, and assigns 1–5 quantile scores mapping to 8 business tiers.")
    ]
    for f_lbl, f_desc in fr_items:
        p = tf_fr.add_paragraph()
        p.space_before = Pt(4)
        r0 = p.add_run()
        r0.text = f_lbl
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_NAVY
        r1 = p.add_run()
        r1.text = f_desc
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_DARK

    # Right Column: Non-Functional + System Specs (Width 4.6")
    add_card(slide3, Inches(7.9), Inches(1.5), Inches(4.633), Inches(2.5))
    tb_nfr = slide3.shapes.add_textbox(Inches(8.1), Inches(1.65), Inches(4.2), Inches(2.2))
    tf_nfr = tb_nfr.text_frame
    tf_nfr.word_wrap = True
    p = tf_nfr.paragraphs[0]
    p.text = "Non-Functional Requirements (NFR)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    nfr_items = [
        ("NFR-01 (Latency): ", "Processes 15,000+ orders in <5.0 seconds using vectorized Pandas C-routines."),
        ("NFR-02 (Reproducibility): ", "Fixed seed (seed=42) guarantees 100% deterministic clustering across platforms."),
        ("NFR-03 (Modularity): ", "Clean star schema data marts (FACT_TRANSACTIONS, DIM_CUSTOMERS).")
    ]
    for n_lbl, n_desc in nfr_items:
        p = tf_nfr.add_paragraph()
        p.space_before = Pt(4)
        r0 = p.add_run()
        r0.text = n_lbl
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_NAVY
        r1 = p.add_run()
        r1.text = n_desc
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_DARK

    # Table 1.1 Summary Card
    add_card(slide3, Inches(7.9), Inches(4.2), Inches(4.633), Inches(2.5))
    tb_tbl = slide3.shapes.add_textbox(Inches(8.1), Inches(4.35), Inches(4.2), Inches(2.2))
    tf_tbl = tb_tbl.text_frame
    tf_tbl.word_wrap = True
    p = tf_tbl.paragraphs[0]
    p.text = "Table 1.1: System Environment Matrix"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    env_items = [
        ("Hardware / CPU: ", "Apple Silicon M-Series / Intel Core i5/i7 (8-Core+)"),
        ("System Memory: ", "8 GB to 16 GB Unified Memory (Min: 4 GB)"),
        ("Storage Drive: ", "Solid State Drive (NVMe SSD)"),
        ("Operating System: ", "macOS Darwin (Unix POSIX) / Linux / Win10+"),
        ("Runtime & Stack: ", "Python 3.9+, Modern Web Browser, VS Code")
    ]
    for e_lbl, e_val in env_items:
        p = tf_tbl.add_paragraph()
        p.space_before = Pt(3)
        r0 = p.add_run()
        r0.text = e_lbl
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_NAVY
        r1 = p.add_run()
        r1.text = e_val
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_DARK

    # =========================================================================
    # SLIDE 4: 1.3 TOOLS AND TECHNOLOGY USED
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Section 1.3  •  Architecture Stack", "Tools & Technology Stack (Multi-Layer Architecture)", 4)
    add_footer(slide4)

    # 4 Horizontal Architectural Layer Cards
    card_w = Inches(11.733)
    layer_h = Inches(1.15)
    layer_gap = Inches(0.18)
    y_start = Inches(1.5)

    layers = [
        ("1. DATA REFINERY & INGESTION LAYER",
         "Python 3.9+  •  Pandas (v2.3.3)  •  NumPy (v2.0.2)",
         "High-throughput tabular ingestion, schema typing, null elimination, reverse logistics isolation, and vectorized Tukey IQR outlier calculations.",
         COLOR_MAROON, RGBColor(254, 242, 242)),
        ("2. MACHINE LEARNING & STATISTICAL MODELING",
         "Scikit-Learn (v1.6.1)  •  SciPy (v1.13.1)",
         "Log1p transformation for Pareto skew compression, StandardScaler normalization, K-Means clustering with k-means++ centroid initialization, and Silhouette score diagnostics.",
         COLOR_BLUE, RGBColor(239, 246, 255)),
        ("3. DIMENSIONAL STORAGE & STAR SCHEMA",
         "Analytical Star Schema Data Marts  •  CSV / Parquet",
         "Decoupled fact and dimension storage (FACT_TRANSACTIONS, FACT_REVERSALS, DIM_CUSTOMERS) designed for BI compatibility with Tableau and Power BI.",
         COLOR_EMERALD, RGBColor(236, 253, 245)),
        ("4. EXECUTIVE BI DASHBOARD & ACTIVATION",
         "HTML5  •  Tailwind CSS  •  Chart.js 4.x  •  Git & GitHub",
         "Client-side canvas graphing engine rendering revenue donuts, spatial RFM bubble charts, KPI scorecards, and searchable CRM activation grids.",
         COLOR_NAVY, RGBColor(248, 250, 252))
    ]

    for idx, (title, stack, desc, tag_color, bg_col) in enumerate(layers):
        y_pos = y_start + (layer_h + layer_gap) * idx
        add_card(slide4, Inches(0.8), y_pos, card_w, layer_h, bg_color=bg_col, border_color=COLOR_BORDER)

        tb = slide4.shapes.add_textbox(Inches(1.0), y_pos + Inches(0.08), card_w - Inches(0.4), layer_h - Inches(0.16))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        r0 = p0.add_run()
        r0.text = title + "  "
        r0.font.name = 'Arial'
        r0.font.size = Pt(11)
        r0.font.bold = True
        r0.font.color.rgb = tag_color

        r1 = p0.add_run()
        r1.text = "—  " + stack
        r1.font.name = 'Arial'
        r1.font.size = Pt(10.5)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_DARK

        p1 = tf.add_paragraph()
        p1.space_before = Pt(3)
        r2 = p1.add_run()
        r2.text = desc
        r2.font.name = 'Arial'
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_SLATE

    # =========================================================================
    # SLIDE 5: 2.1.1 LEVEL 0 CONTEXT DFD
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Section 2.1  •  System Design Using DFD", "Level 0 Data Flow Diagram (Context Architecture)", 5)
    add_footer(slide5)

    # Left: Embedded Visual Card (Width 6.8")
    add_card(slide5, Inches(0.8), Inches(1.5), Inches(6.8), Inches(5.2))
    df0_img = img_dir / "fig_dfd_level_0.png"
    if df0_img.exists():
        slide5.shapes.add_picture(str(df0_img), Inches(1.0), Inches(1.75), width=Inches(6.4))
    tb_cap = slide5.shapes.add_textbox(Inches(0.8), Inches(6.25), Inches(6.8), Inches(0.35))
    p = tb_cap.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Figure 2.1: Level 0 Context Data Flow Diagram (Macro Boundary)"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_SLATE

    # Right: Technical Walkthrough (Width 4.6")
    add_card(slide5, Inches(7.9), Inches(1.5), Inches(4.633), Inches(5.2))
    tb_desc = slide5.shapes.add_textbox(Inches(8.1), Inches(1.65), Inches(4.2), Inches(4.9))
    tf = tb_desc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Context Boundary & Entity Interfaces"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    df0_points = [
        ("Process 0.0 (Global Boundary): ", "Represents the complete 'E-Commerce Customer Lifecycle & RFM Analytics System' interacting with 4 primary external entities."),
        ("1. OLTP Sales Stream: ", "Provides high-frequency transaction logs containing InvoiceNo, StockCode, Quantity, InvoiceDate, UnitPrice, and CustomerID."),
        ("2. Data Scientist: ", "Supplies IQR outlier tuning multipliers, cluster count hyperparameters (k in [2..6]), and evaluates validation diagnostics."),
        ("3. CRM Marketing Engine: ", "Consumes segmented customer lists ('Champions', 'At Risk Whales') for programmatic email, SMS, and ad audience targeting."),
        ("4. Commercial Executives: ", "Receives interactive executive KPI scorecards, segment revenue distribution, and churn risk summaries.")
    ]
    for d_lbl, d_desc in df0_points:
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        r0 = p.add_run()
        r0.text = "• " + d_lbl
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_NAVY
        r1 = p.add_run()
        r1.text = d_desc
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_DARK

    # =========================================================================
    # SLIDE 6: 2.1.2 LEVEL 1 DATA FLOW DIAGRAM
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Section 2.1  •  Pipeline Architecture", "Level 1 Data Flow Diagram (End-to-End Analytics Pipeline)", 6)
    add_footer(slide6)

    # Left: Embedded Visual Card (Width 6.8")
    add_card(slide6, Inches(0.8), Inches(1.5), Inches(6.8), Inches(5.2))
    df1_img = img_dir / "fig_dfd_level_1.png"
    if df1_img.exists():
        slide6.shapes.add_picture(str(df1_img), Inches(1.0), Inches(1.75), width=Inches(6.4))
    tb_cap = slide6.shapes.add_textbox(Inches(0.8), Inches(6.25), Inches(6.8), Inches(0.35))
    p = tb_cap.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Figure 2.2: Level 1 Data Flow Diagram (Pipeline Architecture)"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_SLATE

    # Right: 5 Sequential Processes & Data Stores (Width 4.6")
    add_card(slide6, Inches(7.9), Inches(1.5), Inches(4.633), Inches(5.2))
    tb_desc = slide6.shapes.add_textbox(Inches(8.1), Inches(1.65), Inches(4.2), Inches(4.9))
    tf = tb_desc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "5 Sequential Pipeline Transformations"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    df1_points = [
        ("Process 1.0 (Ingestion & Schema Typing): ", "Parses raw streams, validates typing, and logs records to D1 (Raw Transaction Log)."),
        ("Process 2.0 (Data Quality & Refinery): ", "Drops guest IDs, routes returns to D2 (Reversals Audit Ledger), and applies Tukey IQR fences."),
        ("Process 3.0 (Vectorized RFM Feature Extraction): ", "Computes Recency, Frequency, Spend, and AOV, writing to D3 (RFM Analytical Base Table)."),
        ("Process 4.0 (ML Clustering & Quantile Scoring): ", "Executes log1p scaling, StandardScaler, K-Means, and 1–5 scoring into D4 (Segmented Customer Mart)."),
        ("Process 5.0 (Star Schema & Dashboard): ", "Publishes BI data marts and powers the interactive executive web dashboard.")
    ]
    for d_lbl, d_desc in df1_points:
        p = tf.add_paragraph()
        p.space_before = Pt(5)
        r0 = p.add_run()
        r0.text = "• " + d_lbl
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_NAVY
        r1 = p.add_run()
        r1.text = d_desc
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_DARK

    # =========================================================================
    # SLIDE 7: 2.1.3 LEVEL 2 DATA FLOW DIAGRAM
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Section 2.1  •  Process Decomposition", "Level 2 Data Flow Diagram (Refinery & ML Decomposition)", 7)
    add_footer(slide7)

    # Left: Embedded Visual Card (Width 6.8")
    add_card(slide7, Inches(0.8), Inches(1.5), Inches(6.8), Inches(5.2))
    df2_img = img_dir / "fig_dfd_level_2.png"
    if df2_img.exists():
        slide7.shapes.add_picture(str(df2_img), Inches(1.0), Inches(1.8), width=Inches(6.4))
    tb_cap = slide7.shapes.add_textbox(Inches(0.8), Inches(6.25), Inches(6.8), Inches(0.35))
    p = tb_cap.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Figure 2.3: Level 2 Data Flow Diagram (Detailed Decomposition)"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_SLATE

    # Right: Algorithmic Breakdown (Width 4.6")
    add_card(slide7, Inches(7.9), Inches(1.5), Inches(4.633), Inches(5.2))
    tb_desc = slide7.shapes.add_textbox(Inches(8.1), Inches(1.65), Inches(4.2), Inches(4.9))
    tf = tb_desc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Algorithmic & Mathematical Decomposition"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    df2_points = [
        ("Sub-Process 2.1 (Null ID Purge): ", "Validates CustomerID existence; drops unauthenticated guest sessions."),
        ("Sub-Process 2.2 (Reversals Filter): ", "Detects cancellation identifiers ('^C' or negative Quantity) and exports to fact_reversals_audit.csv."),
        ("Sub-Process 2.3 (Tukey IQR Capping): ", "Computes [Q1 - 1.5*IQR, Q3 + 1.5*IQR] to eliminate wholesale bulk distributor anomalies."),
        ("Sub-Process 4.1 (Log1p Scaling): ", "Applies z = ln(1 + x) to compress severe right-skewed Pareto customer spending."),
        ("Sub-Process 4.2 (StandardScaler): ", "Transforms features to zero mean and unit variance for distance-based clustering."),
        ("Sub-Process 4.3 (K-Means & Silhouette): ", "Fits K-Means across k in [2..6], diagnosing optimal clusters via Silhouette coefficient.")
    ]
    for d_lbl, d_desc in df2_points:
        p = tf.add_paragraph()
        p.space_before = Pt(4)
        r0 = p.add_run()
        r0.text = "• " + d_lbl
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_NAVY
        r1 = p.add_run()
        r1.text = d_desc
        r1.font.size = Pt(8.8)
        r1.font.color.rgb = COLOR_DARK

    # =========================================================================
    # SLIDE 8: 2.2 UML USE CASE DIAGRAM
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Section 2.2  •  System Capabilities", "UML Use Case Diagram & Actor Responsibilities", 8)
    add_footer(slide8)

    # Left: Embedded Visual Card (Width 5.5")
    add_card(slide8, Inches(0.8), Inches(1.5), Inches(5.5), Inches(5.2))
    uc_img = img_dir / "fig_use_case_diagram.png"
    if uc_img.exists():
        slide8.shapes.add_picture(str(uc_img), Inches(1.05), Inches(1.75), width=Inches(5.0))
    tb_cap = slide8.shapes.add_textbox(Inches(0.8), Inches(6.25), Inches(5.5), Inches(0.35))
    p = tb_cap.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Figure 2.4: UML Use Case Diagram (Actors & Capabilities)"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_SLATE

    # Right: Formal Specification Matrix (Width 5.9")
    add_card(slide8, Inches(6.6), Inches(1.5), Inches(5.933), Inches(5.2))
    tb_desc = slide8.shapes.add_textbox(Inches(6.8), Inches(1.65), Inches(5.5), Inches(4.9))
    tf = tb_desc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Table 2.1: Use Case Specification Matrix (UC-01 to UC-06)"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    uc_items = [
        ("UC-01: Ingest Transaction Stream", "Data Engineer / Admin", "Ingests raw event logs, enforces schema typing, verifies datatypes."),
        ("UC-02: Isolate Reversals & IQR Filter", "Data Engineer / Admin", "Purges null IDs, isolates returns (^C), caps Tukey bulk outliers."),
        ("UC-03: Compute Vectorized RFM", "Analytics Engineer", "Aggregates transactions into customer Recency, Frequency, and Spend."),
        ("UC-04: Train & Tune K-Means Model", "Data Scientist", "Applies log1p + StandardScaler, evaluates Elbow/Silhouette, fits k=4."),
        ("UC-05: Inspect Executive KPI Metrics", "Commercial Executive", "Monitors portfolio health ($65k spend), Champions share (72.5%), at-risk capital."),
        ("UC-06: Export Targeted CRM Audiences", "Marketing Manager", "Filters segments in activation ledger and exports audience lists for campaigns.")
    ]
    for u_id, u_act, u_desc in uc_items:
        p = tf.add_paragraph()
        p.space_before = Pt(4)
        r0 = p.add_run()
        r0.text = u_id + " [" + u_act + "]: "
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_NAVY
        r1 = p.add_run()
        r1.text = u_desc
        r1.font.size = Pt(8.8)
        r1.font.color.rgb = COLOR_DARK

    # =========================================================================
    # SLIDE 9: MILESTONE STATUS & PROJECT ROADMAP
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "Milestone Tracking  •  CE-502 / PRJ-701", "Milestone Evaluation Schedule & Project Roadmap", 9)
    add_footer(slide9)

    # Big Summary Card with Table 2.2
    add_card(slide9, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.2))
    tb_m = slide9.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(11.333), Inches(4.9))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True

    p = tf_m.paragraphs[0]
    p.text = "Formal Assessment Schedule & Deliverable Tracking"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    milestones = [
        ("24/09/2026", "First Progress Work Evaluation", "Sections 1.1, 1.2, 1.3, 2.1, 2.2", "10 Marks", "COMPLETED (Current Submission)", COLOR_EMERALD, True),
        ("13/10/2026", "Mid-Term Examination (Second Progress)", "2.3 (Class Diagram), 2.4 (Sequence Diagrams), 3.1 (Data Dictionary), 4.1 (UI Layouts)", "20 Marks", "Scheduled for Next Milestone", COLOR_BLUE, False),
        ("27/10/2026", "Third Progress Work Evaluation", "4.2 (Sample Coding), 4.3 (Navigation), 5.1-5.2 (Conclusion), 6.1 (References)", "10 Marks", "Scheduled for Third Milestone", COLOR_SLATE, False),
        ("06/11/2026", "Report Submission (Spiral Binding)", "Final Signed Spiral Copy & Project Media", "10 Marks", "Scheduled for Final Submission", COLOR_SLATE, False),
        ("18-20/11/2026", "Final Examination (SEE)", "Project Execution Demonstration & Viva Voce", "50 Marks", "Scheduled for Semester Examination", COLOR_MAROON, False)
    ]

    for dt, part, scope, mrk, stat, clr, is_active in milestones:
        p = tf_m.add_paragraph()
        p.space_before = Pt(8)

        r0 = p.add_run()
        r0.text = f"• [{dt}]  {part} ({mrk}) — "
        r0.font.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = clr

        r1 = p.add_run()
        r1.text = stat + "\n"
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = clr

        r2 = p.add_run()
        r2.text = f"   Scope: {scope}"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_DARK

    # Bottom Callout Box
    add_card(slide9, Inches(1.1), Inches(5.6), Inches(11.133), Inches(0.9), bg_color=RGBColor(240, 253, 244), border_color=COLOR_EMERALD)
    tb_call = slide9.shapes.add_textbox(Inches(1.2), Inches(5.65), Inches(10.9), Inches(0.8))
    p = tb_call.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "MILESTONE STATUS SUMMARY (24/09/2026): "
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_EMERALD
    r_txt = p.add_run()
    r_txt.text = "All deliverables for the First Progress Work milestone (10 Marks) have been successfully engineered, verified, and compiled into the 10-page report and interactive demonstration dashboard."
    r_txt.font.size = Pt(9.5)
    r_txt.font.color.rgb = COLOR_DARK

    # =========================================================================
    # SLIDE 10: THANK YOU SLIDE (CONCLUSION & Q&A)
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)

    # Top & Bottom accent bands
    top_bar = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_MAROON
    top_bar.line.color.rgb = COLOR_MAROON

    bot_bar = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.38), Inches(13.333), Inches(0.12))
    bot_bar.fill.solid()
    bot_bar.fill.fore_color.rgb = COLOR_NAVY
    bot_bar.line.color.rgb = COLOR_NAVY

    # IAR Banner
    if iar_banner.exists():
        slide10.shapes.add_picture(str(iar_banner), Inches(3.9), Inches(0.5), width=Inches(5.5))

    # Large Thank You Heading
    tb_thx = slide10.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(1.4))
    p_thx = tb_thx.text_frame.paragraphs[0]
    p_thx.alignment = PP_ALIGN.CENTER
    p_thx.text = "THANK YOU!"
    p_thx.font.name = 'Arial'
    p_thx.font.size = Pt(46)
    p_thx.font.bold = True
    p_thx.font.color.rgb = COLOR_MAROON

    p_qa = tb_thx.text_frame.add_paragraph()
    p_qa.alignment = PP_ALIGN.CENTER
    p_qa.text = "Questions & Technical Discussion"
    p_qa.font.name = 'Arial'
    p_qa.font.size = Pt(20)
    p_qa.font.bold = True
    p_qa.font.color.rgb = COLOR_NAVY

    # Center Summary Card
    add_card(slide10, Inches(2.5), Inches(3.4), Inches(8.333), Inches(3.4), bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)
    tb_end = slide10.shapes.add_textbox(Inches(2.8), Inches(3.55), Inches(7.733), Inches(3.1))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True

    p = tf_end.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "E-COMMERCE CUSTOMER LIFECYCLE & RFM SEGMENTATION DASHBOARD"
    p.font.name = 'Arial'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY

    p_div = tf_end.add_paragraph()
    p_div.alignment = PP_ALIGN.CENTER
    p_div.text = "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━"
    p_div.font.size = Pt(8.5)
    p_div.font.color.rgb = COLOR_MAROON

    end_details = [
        ("Candidate: ", "Mr. Deep Koshiya (B.Sc. in AI & Data Science, Semester V)"),
        ("Institution: ", "Institute of Advanced Research (IAR), Gandhinagar"),
        ("Assessment Milestone: ", "First Progress Work (24/09/2026  •  10 Marks)"),
        ("Course & Project: ", "Project - 1 (CE-502 / PRJ-701)"),
        ("GitHub Repository: ", "https://github.com/Deep-k-coder/ecommerce-rfm-analytics"),
        ("Project Status: ", "Refinery Pipeline, RFM Engine, UML/DFD Models & Dashboard Complete")
    ]
    for lbl, val in end_details:
        p = tf_end.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(3)
        r0 = p.add_run()
        r0.text = lbl
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = COLOR_DARK
        r1 = p.add_run()
        r1.text = val
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_SLATE

    # Save PPTX
    out_pptx_path = base_dir / "Deep_Koshiya_ProjectPresentation_Progress1.pptx"
    prs.save(str(out_pptx_path))
    prs.save(str(base_dir / "Deep_Koshiya_ProjectPresentation.pptx"))
    print(f"[SUCCESS] Presentation saved to {out_pptx_path}")

if __name__ == "__main__":
    create_deck()
